"""pptx skill extract_text unit tests."""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from agentos.skills.bundled.pptx.scripts import extract_text


def test_shape_text_and_table_text_extraction(tmp_path: Path) -> None:
    """extract_text extracts paragraph text and multi-paragraph table cells cleanly."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Textbox shape with paragraph text
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Slide Headline"
    p_body = tb.text_frame.add_paragraph()
    p_body.text = "Secondary point"

    # 2. Table with multi-paragraph cell
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2.5), Inches(5), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text_frame.text = "Metric"
    table.cell(0, 1).text_frame.text = "Value"

    c1 = table.cell(1, 0)
    c1.text_frame.text = "Q1 Revenue"
    p_extra = c1.text_frame.add_paragraph()
    p_extra.text = "(USD)"

    c2 = table.cell(1, 1)
    c2.text_frame.text = "$100M"

    # Test helpers directly
    shape_lines = extract_text._shape_text(tb)
    assert shape_lines == ["Slide Headline", "Secondary point"]

    table_lines = extract_text._table_text(table_shape)
    assert table_lines == ["Metric | Value", "Q1 Revenue (USD) | $100M"]

    # Save presentation and test CLI entrypoint
    pptx_path = tmp_path / "deck.pptx"
    prs.save(str(pptx_path))

    # Test main with --json
    stdout_capture = io.StringIO()
    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(sys, "stdout", stdout_capture)
    try:
        ret = extract_text.main([str(pptx_path), "--json"])
        assert ret == 0
    finally:
        monkeypatch.undo()

    payload = json.loads(stdout_capture.getvalue())
    assert len(payload) == 1
    slide_data = payload[0]
    assert slide_data["slide"] == 1
    assert "Slide Headline" in slide_data["text"]
    assert "Secondary point" in slide_data["text"]
    assert "Metric | Value" in slide_data["text"]
    assert "Q1 Revenue (USD) | $100M" in slide_data["text"]


def _deck_with(text: str, tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = text
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def _code_page_stdout(encoding: str) -> io.TextIOWrapper:
    """A piped stdout whose text layer cannot encode CJK or emoji (#1834)."""
    return io.TextIOWrapper(io.BytesIO(), encoding=encoding, newline="")


def _stdout_bytes(stream: io.TextIOWrapper) -> bytes:
    stream.flush()
    return stream.buffer.getvalue()  # type: ignore[attr-defined]


@pytest.mark.parametrize("code_page", ["cp936", "cp1252"])
def test_plain_output_is_utf8_on_a_non_utf8_code_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, code_page: str
) -> None:
    """The agent captures stdout through a pipe, where Python picks the locale
    code page; extracted CJK or emoji text must still arrive as UTF-8 bytes."""
    deck = _deck_with("季度回顾 🎉", tmp_path)
    stdout = _code_page_stdout(code_page)
    monkeypatch.setattr(sys, "stdout", stdout)

    ret = extract_text.main([str(deck)])

    assert ret == 0
    assert _stdout_bytes(stdout) == "--- slide 1 ---\n季度回顾 🎉\n".encode()


def test_json_output_is_utf8_on_a_non_utf8_code_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = _deck_with("季度回顾 🎉", tmp_path)
    stdout = _code_page_stdout("cp1252")
    monkeypatch.setattr(sys, "stdout", stdout)

    ret = extract_text.main([str(deck), "--json"])

    assert ret == 0
    raw = _stdout_bytes(stdout)
    assert raw.endswith(b"\n")
    assert json.loads(raw.decode("utf-8")) == [{"slide": 1, "text": ["季度回顾 🎉"]}]
